"""Servicio para generar reportes Excel de comparación de precios."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from price_manager.services.audit_service import auditar_accion

from price_manager.services.price_alert_service import (
    comparar_item_web,
    construir_indice_productos_internos,
    leer_jsonl,
)


def obtener_fecha_extraccion(
    fecha_extraccion: datetime | str | None = None,
) -> str:
    """Devuelve la fecha de extracción en formato legible."""
    if fecha_extraccion is None:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(fecha_extraccion, datetime):
        return fecha_extraccion.strftime("%Y-%m-%d %H:%M:%S")

    return str(fecha_extraccion)


def generar_registros_reporte_excel(
    ruta_resultados_web: str | Path,
    productos_internos: list[Any] | None = None,
    cotizacion_usd: float = 1000.0,
    fecha_extraccion: datetime | str | None = None,
) -> list[dict[str, Any]]:
    """Genera los registros necesarios para el reporte Excel."""
    resultados_web = leer_jsonl(ruta_resultados_web)

    indice_productos = construir_indice_productos_internos(
        productos_internos=productos_internos,
        cotizacion_usd=cotizacion_usd,
    )

    fecha = obtener_fecha_extraccion(fecha_extraccion)
    registros = []

    for item_web in resultados_web:
        comparacion = comparar_item_web(
            item_web=item_web,
            indice_productos=indice_productos,
            diferencia_maxima_porcentaje=0.0,
        )

        if comparacion is None:
            continue

        registros.append(
            {
                "Producto": comparacion["producto_buscado"],
                "Precio interno": comparacion["precio_interno_ars"],
                "Precio web": comparacion["precio_web_ars"],
                "Diferencia": comparacion["diferencia_monto"],
                "Fecha de extracción": fecha,
            }
        )

    return registros


def ajustar_ancho_columnas(hoja) -> None:
    """Ajusta el ancho de columnas según el contenido."""
    for columna in hoja.columns:
        largo_maximo = 0
        letra_columna = get_column_letter(columna[0].column)

        for celda in columna:
            valor = celda.value

            if valor is not None:
                largo_maximo = max(largo_maximo, len(str(valor)))

        hoja.column_dimensions[letra_columna].width = largo_maximo + 3

@auditar_accion("generar_reporte_excel_precios")
def generar_reporte_excel_precios(
    ruta_resultados_web: str | Path,
    ruta_reporte_excel: str | Path,
    productos_internos: list[Any] | None = None,
    cotizacion_usd: float = 1000.0,
    fecha_extraccion: datetime | str | None = None,
) -> Path:
    """Genera un archivo Excel con la comparación de precios.

    Args:
        ruta_resultados_web: Archivo JSONL generado por el scraper.
        ruta_reporte_excel: Ruta del archivo Excel a generar.
        productos_internos: Lista opcional de productos internos.
        cotizacion_usd: Cotización utilizada si algún precio interno está en USD.
        fecha_extraccion: Fecha de extracción. Si no se informa, usa la actual.

    Returns:
        Ruta del archivo Excel generado.
    """
    registros = generar_registros_reporte_excel(
        ruta_resultados_web=ruta_resultados_web,
        productos_internos=productos_internos,
        cotizacion_usd=cotizacion_usd,
        fecha_extraccion=fecha_extraccion,
    )

    ruta_reporte_excel = Path(ruta_reporte_excel)
    ruta_reporte_excel.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Reporte Precios"

    columnas = [
        "Producto",
        "Precio interno",
        "Precio web",
        "Diferencia",
        "Fecha de extracción",
    ]

    hoja.append(columnas)

    for registro in registros:
        hoja.append([registro[columna] for columna in columnas])

    encabezado_fill = PatternFill(
        start_color="D9EAF7",
        end_color="D9EAF7",
        fill_type="solid",
    )

    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = encabezado_fill
        celda.alignment = Alignment(horizontal="center")

    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(horizontal="left")

    for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=4):
        for celda in fila:
            celda.number_format = '"$"#,##0.00'

    ajustar_ancho_columnas(hoja)

    workbook.save(ruta_reporte_excel)

    return ruta_reporte_excel

@auditar_accion("ejecutar_scraper_y_generar_reporte_excel")
def ejecutar_scraper_y_generar_reporte_excel(
    productos: list[str] | None = None,
    limite_por_busqueda: int = 10,
    scraper_output_path: str = "data/star_computacion_productos.jsonl",
    reporte_excel_path: str = "data/reporte_precios.xlsx",
    cotizacion_usd: float = 1000.0,
) -> dict[str, Any]:
    """Ejecuta el scraper y genera el reporte Excel."""
    from price_manager.scraper.runner import ejecutar_scraper_star_computacion

    ruta_resultados = ejecutar_scraper_star_computacion(
        productos=productos,
        limite_por_busqueda=limite_por_busqueda,
        output_path=scraper_output_path,
    )

    ruta_reporte = generar_reporte_excel_precios(
        ruta_resultados_web=ruta_resultados,
        ruta_reporte_excel=reporte_excel_path,
        cotizacion_usd=cotizacion_usd,
    )

    return {
        "ruta_resultados_web": Path(ruta_resultados),
        "ruta_reporte_excel": ruta_reporte,
    }